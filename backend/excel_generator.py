"""
Excel report generator for invoice routing data.

Produces a multi-sheet workbook:
  Sheet 1 — Routing Decisions (all invoices processed)
  Sheet 2 — Vendor Profiles (learning data)
  Sheet 3 — Entity Summary (by entity, with totals)
  Sheet 4 — Review Queue (pending + resolved)
"""

from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from models import EntityProfile, ReviewQueue, RoutingDecision, VendorProfile

# ─── Style constants ──────────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
ALT_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
BORDER_SIDE = Side(border_style="thin", color="BFBFBF")
CELL_BORDER = Border(
    left=BORDER_SIDE, right=BORDER_SIDE, top=BORDER_SIDE, bottom=BORDER_SIDE
)

STATUS_COLORS = {
    "auto_routed": "C6EFCE",
    "confirmed": "C6EFCE",
    "pending_review": "FFEB9C",
    "escalated": "FFC7CE",
    "rejected": "FFC7CE",
}

VENDOR_STATUS_COLORS = {
    "new": "D9D9D9",
    "learning": "FFEB9C",
    "confirmed": "BDD7EE",
    "auto_route": "C6EFCE",
}


def _write_headers(ws, headers: list[str], row: int = 1) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER


def _auto_width(ws, max_col: int, min_width: int = 10, max_width: int = 50) -> None:
    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col):
            for cell in row:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 2, max_width))


def _fmt_dt(dt: Optional[datetime]) -> str:
    if dt is None:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M")


def _safe_json_list(raw: Optional[str]) -> str:
    if not raw:
        return ""
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return "; ".join(str(x) for x in parsed)
        return str(parsed)
    except Exception:
        return raw or ""


# ─── Sheet builders ───────────────────────────────────────────────────────────


def _build_decisions_sheet(ws, decisions: list[RoutingDecision]) -> None:
    ws.title = "Routing Decisions"
    ws.freeze_panes = "A2"

    headers = [
        "Invoice #", "Filename", "Vendor (Raw)", "Vendor (Canonical)",
        "Entity", "Vertical", "Confidence", "Tier", "Status",
        "Amount", "Currency", "Invoice Date", "PO Number",
        "Signals Matched", "LLM Reasoning", "Confirmed By", "Confirmed At", "Created At",
    ]
    _write_headers(ws, headers)

    for row_idx, d in enumerate(decisions, start=2):
        fill_color = STATUS_COLORS.get(d.routing_status or "", "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid") if fill_color != "FFFFFF" else None

        values = [
            d.invoice_number or "",
            d.invoice_filename or "",
            d.vendor_name_raw or "",
            d.vendor_canonical or "",
            d.entity_name or "",
            d.entity_vertical or "",
            round(d.confidence_score or 0, 3),
            d.tier_used or "",
            d.routing_status or "",
            d.amount or "",
            d.currency or "USD",
            d.invoice_date or "",
            d.po_number or "",
            _safe_json_list(d.signals_matched),
            d.llm_reasoning or "",
            d.human_confirmed_by or "",
            _fmt_dt(d.human_confirmed_at),
            _fmt_dt(d.created_at),
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(wrap_text=False, vertical="top")
            if row_fill:
                cell.fill = row_fill
            elif row_idx % 2 == 0:
                cell.fill = ALT_FILL

    _auto_width(ws, len(headers))
    ws.row_dimensions[1].height = 30


def _build_vendors_sheet(ws, vendors: list[VendorProfile]) -> None:
    ws.title = "Vendor Profiles"
    ws.freeze_panes = "A2"

    headers = [
        "Canonical Name", "Raw Names", "Category", "Cat. Confidence",
        "Status", "Invoice Count", "Confirmation Count",
        "Confirmed Entity", "Avg Amount", "Min Amount", "Max Amount",
        "First Seen", "Last Seen",
    ]
    _write_headers(ws, headers)

    for row_idx, v in enumerate(vendors, start=2):
        fill_color = VENDOR_STATUS_COLORS.get(v.status or "new", "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        values = [
            v.canonical_name or "",
            _safe_json_list(v.raw_names),
            v.vendor_category or "",
            v.vendor_category_confidence or "",
            v.status or "",
            v.invoice_count or 0,
            v.confirmation_count or 0,
            v.confirmed_entity_name or "",
            round(v.typical_amount_avg or 0, 2),
            round(v.typical_amount_min or 0, 2),
            round(v.typical_amount_max or 0, 2),
            _fmt_dt(v.first_seen),
            _fmt_dt(v.last_seen),
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="top")
            cell.fill = row_fill

    _auto_width(ws, len(headers))
    ws.row_dimensions[1].height = 30


def _build_entities_sheet(ws, entities: list[EntityProfile], decisions: list[RoutingDecision]) -> None:
    ws.title = "Entity Summary"
    ws.freeze_panes = "A2"

    headers = [
        "Entity ID", "Entity Name", "Vertical",
        "Total Invoices Routed", "Total Amount Routed",
        "Auto Routed", "Human Confirmed", "Pending Review",
        "Routing Email", "Review Contact",
    ]
    _write_headers(ws, headers)

    # Build decision counts by entity
    by_entity: dict[str, dict] = {}
    for d in decisions:
        eid = d.entity_id or ""
        if eid not in by_entity:
            by_entity[eid] = {"auto": 0, "confirmed": 0, "pending": 0, "total": 0, "amount": 0}
        by_entity[eid]["total"] += 1
        by_entity[eid]["amount"] += d.amount or 0
        status = d.routing_status or ""
        if status == "auto_routed":
            by_entity[eid]["auto"] += 1
        elif status == "confirmed":
            by_entity[eid]["confirmed"] += 1
        elif status in ("pending_review", "escalated"):
            by_entity[eid]["pending"] += 1

    for row_idx, e in enumerate(entities, start=2):
        counts = by_entity.get(e.id, {"auto": 0, "confirmed": 0, "pending": 0, "total": 0, "amount": 0})
        fill = ALT_FILL if row_idx % 2 == 0 else None

        values = [
            e.id,
            e.name,
            e.vertical,
            counts["total"],
            round(counts["amount"], 2),
            counts["auto"],
            counts["confirmed"],
            counts["pending"],
            e.routing_email or "",
            e.review_contact or "",
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="top")
            if fill:
                cell.fill = fill

    _auto_width(ws, len(headers))
    ws.row_dimensions[1].height = 30


def _build_review_sheet(ws, reviews: list[ReviewQueue]) -> None:
    ws.title = "Review Queue"
    ws.freeze_panes = "A2"

    headers = [
        "Review ID", "Filename", "Vendor", "Amount",
        "Priority", "Status", "Created At", "Resolved At", "Resolved By",
        "Top Candidates",
    ]
    _write_headers(ws, headers)

    for row_idx, r in enumerate(reviews, start=2):
        fill_color = {
            "pending": "FFEB9C",
            "confirmed": "C6EFCE",
            "rejected": "FFC7CE",
        }.get(r.status or "pending", "FFFFFF")
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        candidates_str = ""
        try:
            cands = json.loads(r.top_candidates or "[]")
            candidates_str = "; ".join(
                f"{c.get('entity_name', '')} ({round(c.get('score', 0), 2)})"
                for c in cands[:3]
            )
        except Exception:
            pass

        values = [
            r.id,
            r.invoice_filename or "",
            r.vendor_name or "",
            r.amount or "",
            r.priority or "normal",
            r.status or "pending",
            _fmt_dt(r.created_at),
            _fmt_dt(r.resolved_at),
            r.resolved_by or "",
            candidates_str,
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="top")
            cell.fill = row_fill

    _auto_width(ws, len(headers))
    ws.row_dimensions[1].height = 30


# ─── Public API ───────────────────────────────────────────────────────────────


def generate_excel_report(
    decisions: list[RoutingDecision],
    vendors: list[VendorProfile],
    entities: list[EntityProfile],
    reviews: list[ReviewQueue],
) -> bytes:
    """
    Generate a complete Excel workbook and return it as bytes.
    """
    wb = Workbook()

    # Remove default sheet
    default_ws = wb.active
    wb.remove(default_ws)

    # Add sheets
    ws_decisions = wb.create_sheet("Routing Decisions")
    ws_vendors = wb.create_sheet("Vendor Profiles")
    ws_entities = wb.create_sheet("Entity Summary")
    ws_reviews = wb.create_sheet("Review Queue")

    _build_decisions_sheet(ws_decisions, decisions)
    _build_vendors_sheet(ws_vendors, vendors)
    _build_entities_sheet(ws_entities, entities, decisions)
    _build_review_sheet(ws_reviews, reviews)

    # Metadata sheet
    ws_meta = wb.create_sheet("Report Info")
    ws_meta["A1"] = "Invoice Routing Intelligence — Export"
    ws_meta["A1"].font = Font(bold=True, size=14)
    ws_meta["A3"] = "Generated At"
    ws_meta["B3"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    ws_meta["A4"] = "Total Decisions"
    ws_meta["B4"] = len(decisions)
    ws_meta["A5"] = "Total Vendors"
    ws_meta["B5"] = len(vendors)
    ws_meta["A6"] = "Total Entities"
    ws_meta["B6"] = len(entities)
    ws_meta["A7"] = "Pending Reviews"
    ws_meta["B7"] = sum(1 for r in reviews if r.status == "pending")
    ws_meta.column_dimensions["A"].width = 20
    ws_meta.column_dimensions["B"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
